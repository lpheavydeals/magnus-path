"""Build public/downloads/ads-report-template.xlsx and .csv from docs/ads-report-template-columns.md.

Run: python3 scripts/make-ads-report-template.py
"""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill

OUT = Path(__file__).resolve().parent.parent / 'public' / 'downloads'
NOTE = 'Example numbers only. Replace the blue input cells with your own data. Amounts in your store currency.'
HEAD = ['Month', 'Channel', 'Ad spend', 'Impressions', 'Clicks', 'Orders', 'Revenue', 'New customers', 'Gross margin %',
        'Returning-customer orders', 'CTR', 'CPC', 'Conversion rate', 'AOV', 'ROAS', 'ACoS', 'Cost per order', 'CAC',
        'New-customer share', 'Gross profit', 'Profit after ads', 'Break-even ROAS', 'Notes']
FMT = {'C': '#,##0', 'D': '#,##0', 'E': '#,##0', 'F': '#,##0', 'G': '#,##0', 'H': '#,##0', 'I': '0%', 'J': '#,##0',
       'K': '0.00%', 'L': '#,##0.00', 'M': '0.00%', 'N': '#,##0', 'O': '0.00"x"', 'P': '0.0%', 'Q': '#,##0', 'R': '#,##0',
       'S': '0.0%', 'T': '#,##0', 'U': '#,##0', 'V': '0.00"x"'}
ROWS = {  # row: (month, channel, spend, impressions, clicks, orders, revenue, new customers, margin, notes)
    3: ('Sep 2026', 'Meta ads', 60000, 900000, 12000, 150, 210000, 110, 0.45, 'Prospecting + retargeting'),
    4: ('Sep 2026', 'Google ads', 40000, 300000, 8000, 120, 180000, 70, 0.45, 'Search, Shopping, PMax'),
    5: ('Sep 2026', 'Amazon ads', 30000, 400000, 6000, 140, 140000, 90, 0.30, 'Sponsored Products + Brands; new customers = new-to-brand orders'),
    8: ('Oct 2026', 'Meta ads', 80000, 1150000, 15500, 210, 305000, 150, 0.45, 'Festive creatives from 1 Oct'),
    9: ('Oct 2026', 'Google ads', 45000, 320000, 8600, 140, 215000, 75, 0.45, ''),
    10: ('Oct 2026', 'Amazon ads', 35000, 450000, 6800, 165, 170000, 100, 0.30, ''),
}
TOTALS = {6: ('Sep 2026', 3, 5), 11: ('Oct 2026', 8, 10)}
BLUE, BOLD, ITAL = Font(color='1F4FD1'), Font(bold=True), Font(italic=True, color='555555')


def row_formulas(r):
    return {'J': f'=F{r}-H{r}', 'K': f'=IFERROR(E{r}/D{r},"")', 'L': f'=IFERROR(C{r}/E{r},"")', 'M': f'=IFERROR(F{r}/E{r},"")',
            'N': f'=IFERROR(G{r}/F{r},"")', 'O': f'=IFERROR(G{r}/C{r},"")', 'P': f'=IFERROR(C{r}/G{r},"")',
            'Q': f'=IFERROR(C{r}/F{r},"")', 'R': f'=IFERROR(C{r}/H{r},"")', 'S': f'=IFERROR(H{r}/F{r},"")',
            'T': f'=G{r}*I{r}', 'U': f'=T{r}-C{r}', 'V': f'=IFERROR(1/I{r},"")'}


def build_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Channels'
    ws['A1'] = NOTE
    ws['A1'].font = ITAL
    ws.merge_cells('A1:W1')
    for i, h in enumerate(HEAD, 1):
        c = ws.cell(2, i, h)
        c.font = BOLD
        c.fill = PatternFill('solid', fgColor='F3F1EE')
        c.alignment = Alignment(wrap_text=True, vertical='top')
    for r, v in ROWS.items():
        for col, val in zip('ABCDEFGHI', v[:9]):
            ws[f'{col}{r}'] = val
            ws[f'{col}{r}'].font = BLUE
        ws[f'W{r}'] = v[9]
        ws[f'W{r}'].font = BLUE
        for col, f in row_formulas(r).items():
            ws[f'{col}{r}'] = f
    for r, (month, a, b) in TOTALS.items():
        ws[f'A{r}'], ws[f'B{r}'], ws[f'W{r}'] = month, 'Total', 'Month total'
        for col in 'CDEFGH':
            ws[f'{col}{r}'] = f'=SUM({col}{a}:{col}{b})'
        ws[f'I{r}'] = f'=IFERROR(T{r}/G{r},"")'
        for col, f in row_formulas(r).items():
            ws[f'{col}{r}'] = f
        ws[f'T{r}'] = f'=SUM(T{a}:T{b})'
        for col in 'ABCDEFGHIJKLMNOPQRSTUVW':
            ws[f'{col}{r}'].font = BOLD
    for r in list(ROWS) + list(TOTALS):
        for col, f in FMT.items():
            ws[f'{col}{r}'].number_format = f
    ws.conditional_formatting.add('U3:U11', CellIsRule(operator='lessThan', formula=['0'], font=Font(color='C0392B')))
    widths = {'A': 11, 'B': 13, 'W': 46}
    for i, h in enumerate(HEAD, 1):
        col = ws.cell(2, i).column_letter
        ws.column_dimensions[col].width = widths.get(col, max(11, min(16, len(h) + 2)))
    ws.row_dimensions[2].height = 32
    ws.freeze_panes = 'C3'

    sm = wb.create_sheet('Summary')
    sm['A1'] = NOTE
    sm['A1'].font = ITAL
    sm.merge_cells('A1:K1')
    sh = ['Month', 'Total ad spend', 'Ad-reported revenue', 'Store revenue', 'Store orders', 'Store new customers',
          'MER (blended ROAS)', 'Blended CAC', 'Ad spend % of revenue', 'Store AOV', 'Notes']
    for i, h in enumerate(sh, 1):
        c = sm.cell(2, i, h)
        c.font = BOLD
        c.fill = PatternFill('solid', fgColor='F3F1EE')
        c.alignment = Alignment(wrap_text=True, vertical='top')
    for r, (month, tot, rev, orders, new, note) in {3: ('Sep 2026', 6, 720000, 520, 300, 'Includes organic, email and direct sales'),
                                                     4: ('Oct 2026', 11, 910000, 640, 370, 'Festive month')}.items():
        sm[f'A{r}'] = month
        sm[f'B{r}'] = f'=Channels!C{tot}'
        sm[f'C{r}'] = f'=Channels!G{tot}'
        sm[f'D{r}'], sm[f'E{r}'], sm[f'F{r}'], sm[f'K{r}'] = rev, orders, new, note
        for col in 'ADEFK':
            sm[f'{col}{r}'].font = BLUE
        sm[f'G{r}'] = f'=IFERROR(D{r}/B{r},"")'
        sm[f'H{r}'] = f'=IFERROR(B{r}/F{r},"")'
        sm[f'I{r}'] = f'=IFERROR(B{r}/D{r},"")'
        sm[f'J{r}'] = f'=IFERROR(D{r}/E{r},"")'
        for col, f in {'B': '#,##0', 'C': '#,##0', 'D': '#,##0', 'E': '#,##0', 'F': '#,##0', 'G': '0.00"x"', 'H': '#,##0', 'I': '0.0%', 'J': '#,##0'}.items():
            sm[f'{col}{r}'].number_format = f
    for col, w in zip('ABCDEFGHIJK', [11, 14, 16, 14, 12, 16, 16, 12, 18, 11, 42]):
        sm.column_dimensions[col].width = w
    sm.row_dimensions[2].height = 32

    hw = wb.create_sheet('How to use')
    hw.column_dimensions['A'].width = 110
    lines = ['How to use this ads report template', '',
             '1. Fill one row per channel per month, straight from each ad platform\'s report for the calendar month.',
             '2. Revenue and orders are what the platform reports (its own attribution). Note the attribution window in Notes.',
             '3. New customers: Meta and Google from your store or CRM where possible; Amazon "new-to-brand" orders.',
             '4. Gross margin %: after product cost, shipping, packaging, payment and marketplace fees, before ads.',
             '5. Fill the Summary sheet from your store\'s own reports; MER is the number to trust when channels overlap.',
             '6. Compare each month with the one before, not with someone else\'s numbers.', '',
             'Blue cells are inputs. Black cells are formulas. Total rows recompute ratios from summed totals, never averages.',
             'Need help reading your numbers? Request a free audit: https://www.magnuspath.com/free-ecommerce-audit/']
    for i, t in enumerate(lines, 1):
        hw[f'A{i}'] = t
    hw['A1'].font = Font(bold=True, size=14)
    wb.save(OUT / 'ads-report-template.xlsx')


def build_csv():
    def vals(v):
        sp, im, cl, od, rv, nc, m = v
        d = lambda a, b: a / b if b else ''
        gp = rv * m
        return [od - nc, d(cl, im), d(sp, cl), d(od, cl), d(rv, od), d(rv, sp), d(sp, rv), d(sp, od), d(sp, nc), d(nc, od), gp, gp - sp, d(1, m)]
    rows = []
    for r in [3, 4, 5, 6, 8, 9, 10, 11]:
        if r in ROWS:
            v = ROWS[r]
            inputs, note = list(v[2:9]), v[9]
            month, chan = v[0], v[1]
        else:
            month, a, b = TOTALS[r]
            parts = [ROWS[x] for x in range(a, b + 1)]
            sums = [sum(p[i] for p in parts) for i in range(2, 8)]
            gp = sum(p[6] * p[8] for p in parts)
            inputs, note, chan = sums + [gp / sums[4]], 'Month total', 'Total'
        comp = vals(inputs)
        if r in TOTALS:
            comp[10] = gp
            comp[11] = gp - inputs[0]
        rows.append((month, chan, inputs, comp, note))
    kinds = ['n', 'p', 'm', 'p', 'm', 'r', 'p', 'm', 'm', 'p', 'm', 'm', 'r']
    with open(OUT / 'ads-report-template.csv', 'w', newline='', encoding='utf-8-sig') as fh:
        w = csv.writer(fh)
        w.writerow(HEAD)
        for month, chan, inputs, comp, note in rows:
            out = [month, chan] + [int(x) for x in inputs[:6]] + [round(inputs[6], 4)]
            for k, x in zip(kinds, comp):
                out.append('' if x == '' else int(x) if k == 'n' else round(x, 4) if k == 'p' else round(x, 2))
            w.writerow(out + [note])


if __name__ == '__main__':
    OUT.mkdir(parents=True, exist_ok=True)
    build_xlsx()
    build_csv()
    print('wrote', *sorted(p.name for p in OUT.iterdir()))
